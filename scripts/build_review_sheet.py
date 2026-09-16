#!/usr/bin/env python3
"""Build a spreadsheet of every trial the search found, for expert review.

    python scripts/build_review_sheet.py                    # live fetch
    python scripts/build_review_sheet.py --raw cache.json   # reuse a saved fetch

Nothing is removed: COVID-19 trials, ex vivo cell therapies, unknown-status
trials and unclassifiable ones all get a row, with a suggested action and the
reasons for it. A reviewer records their own call in the "Reviewer decision"
column.
"""
import argparse
import collections
import datetime
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

from trials import classify, config, pipeline  # noqa: E402

ORDER = {"Include": 0, "Review": 1, "Exclude": 2}

# (heading, width, value)
COLUMNS = [
    ("NCT ID", 13, lambda r: r["nct"]),
    ("Title", 60, lambda r: r["title"]),
    ("Suggested action", 16, lambda r: r["action"]),
    ("Why flagged", 76, lambda r: r["why"]),
    ("Flags", 26, lambda r: ", ".join(r["flags"])),
    ("Monitor modality", 17, lambda r: r["modality"]),
    ("Proposed modality", 30, lambda r: r["proposed"]),
    ("Assigned from", 14, lambda r: r["source"]),
    ("Monitor delivery", 17, lambda r: ", ".join(r["delivery"])),
    ("Monitor disease area", 18, lambda r: ", ".join(r["disease_area"])),
    ("Monitor targets", 15, lambda r: ", ".join(r["targets"])),
    ("Monitor topics", 24, lambda r: ", ".join(r["topics"])),
    ("Ex vivo cell therapy", 17, lambda r: "Yes" if r["ex_vivo"] else ""),
    ("COVID-19", 10, lambda r: "Yes" if r["covid"] else ""),
    ("Phase", 19, lambda r: r["phase"]),
    ("Status", 20, lambda r: r["status"]),
    ("Start date", 12, lambda r: r["start"]),
    ("Primary completion", 17, lambda r: r["primary_completion"]),
    ("Enrollment", 11, lambda r: r["enrollment"]),
    ("Lead sponsor", 34, lambda r: r["sponsor"]),
    ("Sponsor type", 13, lambda r: r["sponsor_type"]),
    ("Conditions", 38, lambda r: "; ".join(r["conditions"][:6])),
    ("Interventions", 38, lambda r: classify.intervention_text(r["interventions"])),
    ("Countries", 22, lambda r: ", ".join(r["countries"])),
    ("Matched search terms", 38, lambda r: ", ".join(r["matched_terms"][:8])),
    ("Reviewer decision", 15, lambda r: ""),
    ("Reviewer notes", 34, lambda r: ""),
    ("URL", 36, lambda r: r["url"]),
]
WRAP = {"Title", "Why flagged", "Conditions", "Interventions", "Matched search terms",
        "Proposed modality"}


def fill(hex_):
    return PatternFill("solid", fgColor=hex_)


def trials_sheet(wb, recs):
    ws = wb.create_sheet("Trials")
    heads = [c[0] for c in COLUMNS]
    ws.append(heads)
    for i in range(1, len(heads) + 1):
        cell = ws.cell(1, i)
        cell.fill = fill("1F2A5A")
        cell.font = Font(color="FFFFFF", bold=True, size=10)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30

    col = {h: i + 1 for i, h in enumerate(heads)}
    action_fill = {"Include": fill("E3F2E8"), "Review": fill("FDF3DC"), "Exclude": fill("F7E4E5")}
    for r in recs:
        ws.append([get(r) for _, _, get in COLUMNS])
        n = ws.max_row
        a = ws.cell(n, col["Suggested action"])
        a.fill = action_fill[r["action"]]
        a.font = Font(bold=True, size=10)
        if r["proposed"]:
            ws.cell(n, col["Proposed modality"]).fill = fill("EDE7F6")
        ws.cell(n, col["Reviewer decision"]).fill = fill("EAF0FB")
        u = ws.cell(n, col["URL"])
        u.hyperlink = r["url"]
        u.font = Font(color="1155CC", underline="single", size=10)

    for heading, width, _ in COLUMNS:
        letter = get_column_letter(col[heading])
        ws.column_dimensions[letter].width = width
        if heading in WRAP:
            for n in range(2, ws.max_row + 1):
                ws.cell(n, col[heading]).alignment = Alignment(wrap_text=True, vertical="top")

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(heads))}{ws.max_row}"
    choice = DataValidation(type="list", formula1='"Include,Exclude,Unsure"', allow_blank=True)
    ws.add_data_validation(choice)
    d = get_column_letter(col["Reviewer decision"])
    choice.add(f"{d}2:{d}{ws.max_row}")


def summary_sheet(wb, recs, n_terms):
    sm = wb.create_sheet("Summary")
    sm.append(["RNA clinical trials — scan for expert review"])
    sm.cell(1, 1).font = Font(bold=True, size=14)
    for k, v in [("Source", "ClinicalTrials.gov API v2, interventional studies only"),
                 ("Retrieved", datetime.date.today().isoformat()),
                 ("Search terms", f"{n_terms}, including the RNA Therapeutics Monitor's vocabulary"),
                 ("Total records", len(recs)),
                 ("Nothing removed", "every trial any search term returned has a row")]:
        sm.append([k, v])

    def block(title, counts):
        sm.append([])
        sm.append([title])
        sm.cell(sm.max_row, 1).font = Font(bold=True, size=12)
        for k, v in counts:
            sm.append([k, v, f"{v / len(recs) * 100:.1f}%"])

    block("Suggested action", collections.Counter(r["action"] for r in recs).most_common())
    block("Monitor modality (their vocabulary)",
          collections.Counter(r["modality"] for r in recs if r["modality"]).most_common())
    block("Proposed modality (outside the Monitor's set)",
          collections.Counter(r["proposed"] for r in recs if r["proposed"]).most_common())
    block("Phase", collections.Counter(r["phase"] for r in recs).most_common())
    block("Status", collections.Counter(r["status"] for r in recs).most_common())
    sm.column_dimensions["A"].width = 46
    sm.column_dimensions["B"].width = 56
    sm.column_dimensions["C"].width = 10


README = [
    ("How to use this sheet", ""),
    ("", "Every trial ClinicalTrials.gov returned for the search terms has a row — nothing has "
         "been removed, including COVID-19 trials, ex vivo cell therapies, trials with an unknown "
         "recruitment status, and trials that could not be classified. Sort or filter on "
         "'Suggested action', read 'Why flagged', and record your call in 'Reviewer decision' "
         "with reasoning in 'Reviewer notes'."),
    ("", ""),
    ("Suggested action", "A recommendation, not a decision. Include = a named RNA drug appears in "
                         "the intervention. Review = plausibly in scope but needs a human. Exclude = "
                         "probably noise, with the reason in 'Why flagged'. Only Include rows are "
                         "published."),
    ("Monitor modality", "The RNA Therapeutics Monitor's own modality vocabulary, used verbatim: "
                         "mRNA, saRNA, circRNA, siRNA, ASO, aptamer, miRNA, RNA editing, CRISPR RNA, "
                         "RNA nanostructure. The published list uses these labels and no others."),
    ("Proposed modality", "Shaded purple. Where a trial looks real but the Monitor's labels cannot "
                          "hold it, a category outside their vocabulary is proposed rather than "
                          "discarding the trial — DNAzyme, ex vivo cell therapy, exosome delivery, "
                          "and so on. These are for a reviewer to accept or reject; none reach the "
                          "published list unless the vocabulary is deliberately extended."),
    ("Assigned from", "Which layer made the call. 'intervention' is strongest; 'title' and "
                      "'sponsor' are weaker and are said so in 'Why flagged'."),
    ("Monitor * columns", "The Monitor's other category patterns applied to each trial. Coverage is "
                          "poor on trial records because those patterns are tuned for papers, where "
                          "an abstract discusses mechanism. A trial registration names a condition "
                          "and an intervention, not a mechanism. They inform; they decide nothing."),
    ("", ""),
    ("Known limits", ""),
    ("", "1. ClinicalTrials.gov ORs the words of an unquoted multi-word query, so 'nucleic acid "
         "medicine' returns about 2,000 trials while the quoted phrase returns none. Records that "
         "arrived only that way are flagged 'loose-term-only'."),
    ("", "2. Early-phase RNA drugs are often named only by sponsor code (OLX10010, AZD9150, CDR132L) "
         "with no RNA word in the record. The known code series are recovered, and that list is "
         "certainly incomplete — a missing company is the most valuable correction a reviewer can "
         "make."),
    ("", "3. DNAzymes (e.g. SB010, SB011, SB012) are catalytic DNA, not RNA. They are kept with a "
         "proposed label because they are oligonucleotide therapeutics; whether they belong is a "
         "scope decision, not a data one."),
    ("", "4. No relevance model has been applied. Classification is deterministic pattern matching, "
         "documented in trials/classify.py."),
]


def readme_sheet(wb):
    rm = wb.create_sheet("Read me", 0)
    for a, b in README:
        rm.append([a, b])
        if a and not b:
            rm.cell(rm.max_row, 1).font = Font(bold=True, size=13)
        elif a:
            rm.cell(rm.max_row, 1).font = Font(bold=True, size=10)
        rm.cell(rm.max_row, 2).alignment = Alignment(wrap_text=True, vertical="top")
    rm.column_dimensions["A"].width = 24
    rm.column_dimensions["B"].width = 108
    for n in range(1, rm.max_row + 1):
        v = rm.cell(n, 2).value
        if v and len(str(v)) > 100:
            rm.row_dimensions[n].height = 15 * (len(str(v)) // 100 + 1)


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--raw", help="read a saved fetch instead of querying the API")
    ap.add_argument("--out", default="rna-clinical-trials-review.xlsx")
    args = ap.parse_args()

    studies, matched = pipeline.load_or_fetch(args.raw)
    recs = pipeline.records(studies, matched)
    recs.sort(key=lambda r: (ORDER[r["action"]], r["modality"] or r["proposed"] or "zz",
                             r["start"] or ""))

    wb = Workbook()
    wb.remove(wb.active)
    trials_sheet(wb, recs)
    summary_sheet(wb, recs, len(config.search_terms()[0]))
    readme_sheet(wb)
    wb.save(args.out)
    print(f"{args.out}: {len(recs)} trials", file=sys.stderr)


if __name__ == "__main__":
    main()
